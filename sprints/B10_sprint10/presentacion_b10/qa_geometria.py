"""qa_geometria.py — QA sin mirar: tinta por renglón y colisiones dentro de los grupos.

Lo escribió la sesión que construyó el deck cuando no pudo abrir imágenes, como sustituto parcial
de mirar las láminas. NO lo reemplaza: no ve colores, no ve qué dice el texto y no ve la
proximidad entre dos formas que no se tocan ([[deck-qa-puntos-ciegos-chequeo]]).

Por lámina imprime las bandas de tinta del rasterizado (renglones con píxeles < 200: cada banda es
un renglón o un objeto, y los huecos son el aire real) y, dentro de los group shapes, avisa de:
  - texto contra texto, con la extensión REAL del texto (ancho medido con Barlow y alineación),
    no la caja;
  - recta o tramo de polilínea que cruza un texto, salvo que un rectángulo blanco pintado después
    lo tape (la leyenda interna de O1);
  - forma contra texto, y forma a menos de 0,03" de un texto.

Uso, después de rasterizar a <dir>/<prefijo>-N.png con pdftoppm:
    PYTHONPATH=/media/administrador/Storage1/sdonoso/clam_testing2/.pylibs \
      /home/sdonoso/miniconda3/envs/pruebas/bin/python qa_geometria.py <dir> <pptx> <prefijo>
"""
import math, sys
import numpy as np
from PIL import Image
sys.path.insert(0, "/media/administrador/Storage1/sdonoso/clam_testing2/oncomets-ernesto/sprints/B9_sprint9/presentacion_b9")
from generate_b9_deck import text_w
from pptx import Presentation
from pptx.util import Emu
from pptx.enum.shapes import MSO_SHAPE_TYPE, MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN

SP = sys.argv[1]; PPTX = sys.argv[2]; PREF = sys.argv[3]; DPI = 110
E = lambda v: Emu(v).inches

def bandas(png, x0=0.0, x1=13.333):
    a = np.array(Image.open(png).convert("L"))
    a = a[:, int(x0 * DPI):int(x1 * DPI)]
    ink = (a < 200).sum(axis=1)
    rows = np.where(ink > 0)[0]
    out = []
    if len(rows):
        s = p = rows[0]
        for r in rows[1:]:
            if r != p + 1:
                out.append((s, p)); s = r
            p = r
        out.append((s, p))
    return [(s / DPI, (e + 1) / DPI) for s, e in out]

def hijos(shapes, z=[0]):
    for sh in shapes:
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from hijos(sh.shapes, z)
        else:
            z[0] += 1
            yield z[0], sh

def es_blanco(sh):
    try:
        return sh.fill.type == 1 and str(sh.fill.fore_color.rgb) == "FFFFFF"
    except Exception:
        return False

def ext_texto(sh):
    tf = sh.text_frame
    p = tf.paragraphs[0]
    r = p.runs[0]
    fs = r.font.size.pt; bold = bool(r.font.bold)
    txt = tf.text
    tw = text_w(txt, fs, bold)
    l, t, w, h = E(sh.left), E(sh.top), E(sh.width), E(sh.height)
    ml = E(tf.margin_left); mr = E(tf.margin_right)
    al = p.alignment
    if al == PP_ALIGN.CENTER:
        a = l + w / 2 - tw / 2
    elif al == PP_ALIGN.RIGHT:
        a = l + w - mr - tw
    else:
        a = l + ml
    c = t + h / 2
    return (a, c - 0.36 * fs / 72, a + tw, c + 0.40 * fs / 72), txt

def seg_rect(p, q, R, pad=0.0):
    (x1, y1), (x2, y2) = p, q
    a, b, c, d = R[0] - pad, R[1] - pad, R[2] + pad, R[3] + pad
    # Liang-Barsky
    dx, dy = x2 - x1, y2 - y1
    t0, t1 = 0.0, 1.0
    for pp, qq in ((-dx, x1 - a), (dx, c - x1), (-dy, y1 - b), (dy, d - y1)):
        if pp == 0:
            if qq < 0: return False
        else:
            r = qq / pp
            if pp < 0: t0 = max(t0, r)
            else: t1 = min(t1, r)
            if t0 > t1: return False
    return True

def inter(A, B, tol=0.0):
    return A[0] < B[2] - tol and B[0] < A[2] - tol and A[1] < B[3] - tol and B[1] < A[3] - tol

def dist(A, B):
    dx = max(0, max(A[0], B[0]) - min(A[2], B[2]))
    dy = max(0, max(A[1], B[1]) - min(A[3], B[3]))
    return math.hypot(dx, dy)

prs = Presentation(PPTX)
for idx, slide in enumerate(prs.slides, start=1):
    png = "%s/%s-%d.png" % (SP, PREF, idx)
    b = bandas(png)
    print("\n== s%02d  tinta por renglón (pulgadas, alto, hueco previo)" % idx)
    prev = None
    lin = []
    for s, e in b:
        lin.append("%.2f-%.2f(%.2f|%s)" % (s, e, e - s, "-" if prev is None else "%.2f" % (s - prev)))
        prev = e
    print("  " + "  ".join(lin))
    if idx == 1:
        continue
    textos, segs, cajas = [], [], []
    for top in slide.shapes:
        if top.shape_type != MSO_SHAPE_TYPE.GROUP:
            continue
        for z, sh in hijos(top.shapes):
            if sh.shape_type == MSO_SHAPE_TYPE.LINE or sh.__class__.__name__ == "Connector":
                segs.append((z, (E(sh.begin_x), E(sh.begin_y)), (E(sh.end_x), E(sh.end_y)), sh.name))
                continue
            l, t, w, h = E(sh.left), E(sh.top), E(sh.width), E(sh.height)
            geom = sh._element.find(".//{http://schemas.openxmlformats.org/drawingml/2006/main}custGeom")
            if geom is not None:
                pts = [(l + E(int(pt.get("x"))), t + E(int(pt.get("y"))))
                       for pt in geom.iter("{http://schemas.openxmlformats.org/drawingml/2006/main}pt")]
                for p, q in zip(pts[:-1], pts[1:]):
                    segs.append((z, p, q, "polilínea"))
                continue
            tiene_txt = sh.has_text_frame and sh.text_frame.text.strip()
            if tiene_txt and getattr(sh, "shape_type", None) == MSO_SHAPE_TYPE.TEXT_BOX:
                R, txt = ext_texto(sh)
                textos.append((z, R, txt))
            else:
                cajas.append((z, (l, t, l + w, t + h), sh, es_blanco(sh),
                              sh.text_frame.text if tiene_txt else ""))
    avisos = []
    for i in range(len(textos)):
        for j in range(i + 1, len(textos)):
            if inter(textos[i][1], textos[j][1], 0.005):
                avisos.append("texto/texto  «%s» × «%s»" % (textos[i][2][:30], textos[j][2][:30]))
    for zs, p, q, nom in segs:
        for zt, R, txt in textos:
            if seg_rect(p, q, R):
                tapado = any(bl and zc > zs and C[0] <= R[0] and C[2] >= R[2] and C[1] <= R[1] and C[3] >= R[3]
                             for zc, C, sh, bl, _ in cajas)
                if not tapado:
                    avisos.append("línea/texto  %s × «%s»" % (nom, txt[:40]))
    for zc, C, sh, bl, own in cajas:
        if bl:
            continue
        for zt, R, txt in textos:
            if inter(C, R, 0.003):
                avisos.append("forma/texto  %s(%s) × «%s»" % (sh.name, own[:12], txt[:40]))
            elif dist(C, R) < 0.03:
                avisos.append("CERCA %.3f  %s × «%s»" % (dist(C, R), sh.name, txt[:40]))
    print("  %d textos, %d segmentos, %d formas en grupos" % (len(textos), len(segs), len(cajas)))
    for a in avisos:
        print("   ·", a)
